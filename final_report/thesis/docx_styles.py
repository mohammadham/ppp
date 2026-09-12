"""Persian Word layout; no rewriting is done by this module."""
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def element(parent, name, **attrs):
    node = OxmlElement('w:' + name)
    for key, value in attrs.items():
        node.set(qn('w:' + key), str(value))
    parent.append(node)
    return node


def direction(p, rtl=True):
    pr = p._p.get_or_add_pPr()
    for old in pr.findall(qn('w:bidi')):
        pr.remove(old)
    element(pr, 'bidi', val='1' if rtl else '0')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if rtl else WD_ALIGN_PARAGRAPH.LEFT


def runtext(p, text, bold=False, size=None):
    # Separate Latin spans to preserve math, DOI, and mixed-direction author names.
    for part in re.split(r'([A-Za-z0-9][A-Za-z0-9_ .,:;()&+×=<>/%^−–—\-\[\]{}\\]*)', text):
        if not part:
            continue
        r = p.add_run(part)
        latin = bool(re.match(r'[A-Za-z0-9]', part))
        r.font.name = 'Times New Roman' if latin else 'B Lotus'
        r.font.size = Pt(size or (12 if latin else 14))
        r.bold = bold
        pr = r._r.get_or_add_rPr()
        fonts = pr.find(qn('w:rFonts'))
        fonts.set(qn('w:cs'), 'B Lotus')
        element(pr, 'rtl', val='0' if latin else '1')
        element(pr, 'szCs', val=str((size or 14) * 2))
        element(pr, 'lang', val='en-US' if latin else 'fa-IR', bidi='fa-IR')
        if bold:
            element(pr, 'bCs', val='1')


def paragraph(doc, text='', style=None, rtl=True, center=False):
    p = doc.add_paragraph(style=style)
    direction(p, rtl)
    p.paragraph_format.widow_control = True
    for i, chunk in enumerate(re.split(r'\*\*(.*?)\*\*', text)):
        runtext(p, chunk.replace('`', ''), bold=bool(i % 2) or bool(style and style.startswith('Heading')))
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
    return p


def field(p, instruction, cached=''):
    r = p.add_run()._r
    element(r, 'fldChar', fldCharType='begin', dirty='true')
    element(r, 'instrText').text = instruction
    element(r, 'fldChar', fldCharType='separate')
    if cached:
        runtext(p, cached)
    element(p.add_run()._r, 'fldChar', fldCharType='end')


def setup():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = s.right_margin = Cm(3)
    s.left_margin = Cm(2)
    s.footer_distance = Cm(1.5)
    normal = d.styles['Normal']
    normal.font.name, normal.font.size = 'B Lotus', Pt(14)
    element(normal.element.get_or_add_rPr(), 'rFonts', ascii='Times New Roman', hAnsi='Times New Roman', cs='B Lotus')
    element(normal.element.get_or_add_rPr(), 'szCs', val=28)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(.5)
    for n in range(1, 4):
        style = d.styles[f'Heading {n}']
        style.font.name, style.font.size = 'B Lotus', Pt(14)
        style.font.bold, style.font.color.rgb = True, RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(16 if n > 1 else 64)
        style.paragraph_format.space_after = Pt(12)
    d.styles['Caption'].font.name = 'B Lotus'
    d.styles['Caption'].font.size = Pt(11)
    element(d.settings.element, 'updateFields', val='true')
    element(d.settings.element, 'themeFontLang', val='en-US', bidi='fa-IR')
    return d


def markdown(doc, text):
    lines = text.strip().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line or line == '---':
            continue
        if line.startswith('|'):
            rows = [line]
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(lines[i].strip()); i += 1
            cells = [[c.strip() for c in row.strip('|').split('|')] for row in rows if not re.match(r'^\|[\s:|\-]+\|$', row)]
            table = doc.add_table(rows=0, cols=len(cells[0]))
            table.style = 'Table Grid'
            element(table._tbl.tblPr, 'bidiVisual', val='1')
            for index, row in enumerate(cells):
                tr = table.add_row()
                element(tr._tr.get_or_add_trPr(), 'cantSplit')
                if index == 0:
                    element(tr._tr.get_or_add_trPr(), 'tblHeader')
                for cell, value in zip(tr.cells, row):
                    p = cell.paragraphs[0]; direction(p)
                    p.paragraph_format.first_line_indent = Cm(0)
                    p.paragraph_format.space_after = Pt(3)
                    runtext(p, value, bold=index == 0, size=11)
            continue
        match = re.match(r'^(#{1,3}) (.+)', line)
        if match:
            paragraph(doc, match[2], f'Heading {len(match[1])}')
        elif line.startswith('$$'):
            paragraph(doc, line.strip('$'), rtl=False, center=True)
        elif line.startswith('جدول ') or line.startswith('شکل '):
            p = paragraph(doc, line, 'Caption', center=True)
            p.paragraph_format.keep_with_next = line.startswith('جدول ')
        else:
            paragraph(doc, line)